import os
import pandas as pd
from openpyxl import Workbook, load_workbook
from src.configuracion_v1 import EXCEL_PATH, CARPETA_REPORTES



def normalizar_texto(texto):
    """Elimina espacios extra y convierte a mayúsculas para evitar falsos positivos en la detección de cambios."""
    if isinstance(texto, str):
        return " ".join(texto.upper().strip().split())
    return texto


def guardar_en_excel(df):
    """Guarda los cambios detectados en un archivo Excel con una estructura más clara y ordenada."""

    if not os.path.exists(CARPETA_REPORTES):
        os.makedirs(CARPETA_REPORTES)

    # Crear un nuevo archivo si no existe
    if os.path.exists(EXCEL_PATH):
        workbook = load_workbook(EXCEL_PATH)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "Fecha", "Número Expediente", "ID Expediente", "Situación Anterior", "Situación Nueva",
            "Última Actuación Anterior", "Última Actuación Nueva", "Carátula Anterior",
            "Carátula en Base de Datos", "Carátula Nueva", "Tipo de Cambio"
        ])

    # **Reformatear la estructura del DataFrame**
    df = df.copy()  # Evitar modificar el original

    # **Asegurar que todas las columnas necesarias existen y no contienen NaN**
    columnas_correctas = [
        "Fecha", "Número Expediente", "ID Expediente", "Situación Anterior", "Situación Nueva",
        "Última Actuación Anterior", "Última Actuación Nueva", "Carátula Anterior",
        "Carátula en Base de Datos", "Carátula Nueva", "Tipo de Cambio"
    ]

    for col in columnas_correctas:
        if col not in df.columns:
            df[col] = "N/A"
    df = df[columnas_correctas]  # **Reordenar las columnas en el orden correcto**
    df.fillna("N/A", inplace=True)

    # **Convertir fechas a string para evitar NaT y manejar errores**
    for col in ["Situación Nueva", "Última Actuación Nueva"]:
        df[col] = df[col].astype(str).replace({"NaT": "N/A", "nan": "N/A"})

    # **Corregir "Carátula en Base de Datos" en expedientes nuevos**
    df.loc[df["ID Expediente"] == "Nuevo", "Carátula en Base de Datos"] = "N/A"

    # **Normalizar las carátulas antes de comparar**
    df["Carátula Anterior"] = df["Carátula Anterior"].apply(normalizar_texto)
    df["Carátula Nueva"] = df["Carátula Nueva"].apply(normalizar_texto)

    # **No recalcular "Tipo de Cambio", usar el valor de "comparar_expedientes()"**
    # **Evitar sobrescribir valores correctos en "Carátula en Base de Datos"**
    df.loc[df["Carátula Anterior"] == df["Carátula Nueva"], ["Carátula Anterior", "Carátula Nueva"]] = "SIN CAMBIO"

    # **Guardar en Excel sin modificar la clasificación de cambios**
    for _, row in df.iterrows():
        sheet.append(row.tolist())

    workbook.save(EXCEL_PATH)
    print(f"📊 Reporte guardado en {EXCEL_PATH}")


# 📌 FUNCIÓN PARA GUARDAR UN REPORTE DE ACTUACIONES EN EXCEL
def guardar_reporte_actuaciones(actuaciones):
    """Guarda las actuaciones insertadas en un archivo Excel para auditoría."""
    carpeta_reportes = CARPETA_REPORTES
    excel_path = os.path.join(carpeta_reportes, "reporte_actuaciones.xlsx")

    if not os.path.exists(carpeta_reportes):
        os.makedirs(carpeta_reportes)

    if os.path.exists(excel_path):
        workbook = load_workbook(excel_path)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Expediente", "Oficina", "Fecha", "Tipo", "Detalle", "Foja", "Archivo"])

    for actuacion in actuaciones:
        sheet.append(actuacion)

    workbook.save(excel_path)
    print(f"📊 Reporte de actuaciones guardado en {excel_path}")
